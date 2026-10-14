from __future__ import annotations

import hashlib
import re
import uuid
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

try:
    import pdfplumber
except ImportError:
    pdfplumber = None
from docx import Document as WordDocument
from pypdf import PdfReader


@dataclass(frozen=True)
class Page:
    number: int
    content: str
    section_title: str = ""


def validate_file(path: Path, extension: str) -> None:
    if extension == ".pdf":
        with path.open("rb") as handle:
            signature = handle.read(5)
        if signature != b"%PDF-":
            raise ValueError("The uploaded file is not a valid PDF")
    elif extension == ".docx":
        if not zipfile.is_zipfile(path):
            raise ValueError("The uploaded file is not a valid DOCX")
        with zipfile.ZipFile(path) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ValueError("The uploaded file is not a valid DOCX")
    elif extension == ".txt":
        path.read_text(encoding="utf-8-sig")
    else:
        raise ValueError(f"Unsupported file type: {extension}")


def parse_document(path: Path, extension: str) -> list[Page]:
    validate_file(path, extension)
    if extension == ".pdf":
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported")
        if pdfplumber is not None:
            with pdfplumber.open(str(path)) as document:
                extracted = [
                    _clean_text(page.extract_text(x_tolerance=2, y_tolerance=3) or "")
                    for page in document.pages
                ]
        else:
            extracted = [
                _clean_text(page.extract_text() or "")
                for page in reader.pages
            ]
        pages = _prepare_pdf_pages(extracted)
    elif extension == ".docx":
        document = WordDocument(str(path))
        blocks: list[str] = []
        current_heading = ""
        for paragraph in document.paragraphs:
            text = _clean_text(paragraph.text)
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                current_heading = text
                blocks.append(text)
            else:
                blocks.append(f"{current_heading}\n{text}" if current_heading and not blocks else text)
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(_clean_text(cell.text) for cell in row.cells)
                if row_text.strip(" |"):
                    blocks.append(row_text)
        pages = [Page(number=1, content="\n\n".join(blocks), section_title=current_heading)]
    else:
        raw = path.read_text(encoding="utf-8-sig")
        pages = [Page(number=index, content=_clean_text(text)) for index, text in enumerate(raw.split("\f"), start=1)]
    usable = [page for page in pages if page.content.strip()]
    if not usable:
        raise ValueError("No readable text was found in the document")
    return usable


def _clean_text(value: str) -> str:
    value = value.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _prepare_pdf_pages(extracted: list[str]) -> list[Page]:
    """Remove repeated page furniture and retain a useful section label."""
    first_lines = [lines[0] for text in extracted if (lines := _content_lines(text))]
    minimum_repetitions = 2 if len(extracted) < 6 else 3
    repeated_headers = {
        line
        for line, count in Counter(first_lines).items()
        if count >= minimum_repetitions
    }

    pages: list[Page] = []
    for number, text in enumerate(extracted, start=1):
        lines = _content_lines(text)
        if lines and lines[0] in repeated_headers:
            lines.pop(0)
        while lines and re.fullmatch(r"(?:page\s*)?\d{1,4}", lines[-1], flags=re.IGNORECASE):
            lines.pop()
        content = "\n".join(lines).strip()
        section_title = lines[0][:160] if lines else ""
        pages.append(Page(number=number, content=content, section_title=section_title))
    return pages


def _content_lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def _token_spans(text: str) -> list[tuple[int, int]]:
    return [(match.start(), match.end()) for match in re.finditer(r"[A-Za-z0-9_]+|[\u3400-\u9fff]|[^\s]", text)]


def chunk_pages(pages: list[Page], max_words: int, overlap_words: int) -> list[dict[str, object]]:
    chunks: list[dict[str, object]] = []
    chunk_index = 0
    for page in pages:
        spans = _token_spans(page.content)
        if not spans:
            continue
        start = 0
        while start < len(spans):
            end = min(start + max_words, len(spans))
            content = page.content[spans[start][0] : spans[end - 1][1]].strip()
            section = page.section_title or next((line.strip() for line in content.splitlines() if line.strip()), "")
            chunks.append(
                {
                    "id": str(uuid.uuid4()),
                    "chunk_index": chunk_index,
                    "page_start": page.number,
                    "page_end": page.number,
                    "section_title": section[:160],
                    "content": content,
                    "content_hash": hashlib.sha256(content.encode("utf-8")).hexdigest(),
                    "embedding": None,
                }
            )
            chunk_index += 1
            if end == len(spans):
                break
            start = max(start + 1, end - overlap_words)
    return chunks
